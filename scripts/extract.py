"""PPTX -> editable JSON + original/optimized assets. Usage: python scripts/extract.py input.pptx"""
from pathlib import Path
from pptx import Presentation
from pptx.oxml.ns import qn
from PIL import Image
import sys,json,hashlib,io,re,shutil
ROOT=Path(__file__).resolve().parents[1]
p=Presentation(sys.argv[1]); U=12700

def n(v): return round(v/U,3)
def color(el,default=None):
 if el is None:return default
 if el.find(qn('a:noFill')) is not None:return 'transparent'
 f=el.find(qn('a:solidFill'))
 if f is None:return default
 c=f.find(qn('a:srgbClr'))
 if c is not None:
  v='#'+c.get('val');a=c.find(qn('a:alpha'))
  return v if a is None else v+f'{round(int(a.get("val"))*255/100000):02x}'
 s=f.find(qn('a:schemeClr'))
 return {'tx1':'#E8EBF0','bg1':'#0A0D14','accent1':'#42665B','accent2':'#526A88'}.get(s.get('val'),default) if s is not None else default

def paragraphs(tf):
 result=[]
 for para in tf.paragraphs:
  pp=para._p.find(qn('a:pPr')); runs=[]
  for child in para._p:
   if child.tag==qn('a:br'):runs.append({'text':'\n'});continue
   if child.tag!=qn('a:r'):continue
   tx=child.find(qn('a:t')); rp=child.find(qn('a:rPr'))
   if tx is None:continue
   default=pp.find(qn('a:defRPr')) if pp is not None else None
   rp=rp if rp is not None else default if default is not None else child
   font=rp.find(qn('a:latin'))
   if font is None and default is not None:font=default.find(qn('a:latin'))
   runs.append({'text':tx.text or '', 'size':int(rp.get('sz','1600'))/100,'font':font.get('typeface','Inter') if font is not None else 'Inter','bold':rp.get('b')=='1','italic':rp.get('i')=='1','color':color(rp,'#E8EBF0'),'spacing':int(rp.get('spc','0'))/100})
  def space(tag):
   if pp is None:return 0
   a=pp.find(qn('a:'+tag));pt=a.find(qn('a:spcPts')) if a is not None else None
   return int(pt.get('val'))/100 if pt is not None else 0
  ls=pp.find(qn('a:lnSpc')) if pp is not None else None
  pct=ls.find(qn('a:spcPct')) if ls is not None else None
  pts=ls.find(qn('a:spcPts')) if ls is not None else None
  bullet=pp.find(qn('a:buChar')) if pp is not None else None
  result.append({'runs':runs,'text':para.text,'align':pp.get('algn','l') if pp is not None else 'l','lineHeight':int(pct.get('val'))/100000 if pct is not None else 1.15,'lineHeightPt':int(pts.get('val'))/100 if pts is not None else None,'before':space('spcBef'),'after':space('spcAft'),'bullet':bullet.get('char') if bullet is not None else None,'marginLeft':n(int(pp.get('marL','0'))) if pp is not None else 0,'indent':n(int(pp.get('indent','0'))) if pp is not None else 0})
 return result

assets={}; original_dir=ROOT/'sources/media';original_dir.mkdir(parents=True,exist_ok=True)
def shape(sh,background=False):
 d={'id':sh.shape_id,'x':n(sh.left),'y':n(sh.top),'w':n(sh.width),'h':n(sh.height),'rotation':sh.rotation,'background':background}
 if sh.shape_type==13:
  blob=sh.image.blob; key=hashlib.sha256(blob).hexdigest()[:14]; ext=sh.image.ext
  if key not in assets:
   orig=original_dir/f'{key}.{ext}';orig.write_bytes(blob)
   im=Image.open(io.BytesIO(blob));w,h=im.size
   if im.mode not in ['RGB','RGBA']:im=im.convert('RGBA')
   variants=[]
   for maxw in sorted(set([min(w,800),min(w,1600),w if w<1600 else 1600])):
    out=im.copy();out.thumbnail((maxw,100000)); name=f'{key}-{out.width}.webp'
    if not (ROOT/'public/assets/images'/name).exists():out.save(ROOT/'public/assets/images'/name,'WEBP',quality=86,method=6)
    variants.append({'src':'assets/images/'+name,'width':out.width,'height':out.height})
   assets[key]={'original':f'sources/media/{key}.{ext}','width':w,'height':h,'variants':variants}
  d.update(type='image',asset=key,**assets[key]);return d
 if sh.has_table:
  t=sh.table;rows=[]
  for row in t.rows:
   cells=[]
   for c in row.cells:
    cp=c._tc.find(qn('a:tcPr'))
    cells.append({'paragraphs':paragraphs(c.text_frame),'fill':color(cp,'transparent'),'padding':[n(c.margin_top),n(c.margin_right),n(c.margin_bottom),n(c.margin_left)]})
   rows.append({'height':n(row.height),'cells':cells})
  d.update(type='table',columns=[n(c.width) for c in t.columns],rows=rows);d['h']=sum(r['height'] for r in rows);return d
 sp=sh._element.find(qn('p:spPr'));geom=sp.find(qn('a:prstGeom')) if sp is not None else None
 ln=sp.find(qn('a:ln')) if sp is not None else None
 d.update(type='shape',fill=color(sp,'transparent'),border=color(ln,'transparent'),borderWidth=n(int(ln.get('w','0'))) if ln is not None else 0,geometry=geom.get('prst','rect') if geom is not None else 'rect')
 if sh.has_text_frame and sh.text.strip():
  tf=sh.text_frame;body=tf._txBody.find(qn('a:bodyPr')); fit=body.find(qn('a:normAutofit'))
  d.update(type='text',paragraphs=paragraphs(tf),wrap=body.get('wrap','square')!='none',padding=[n(tf.margin_top),n(tf.margin_right),n(tf.margin_bottom),n(tf.margin_left)],anchor=body.get('anchor','t'),fontScale=int(fit.get('fontScale','100000'))/100000 if fit is not None else 1)
 return d

slides=[]
for i,s in enumerate(p.slides):
 shapes=[shape(sh,True) for sh in s.slide_layout.shapes]+[shape(sh) for sh in s.shapes]
 texts=[sh for sh in shapes if sh['type']=='text'];title=max(texts,key=lambda sh:max((r.get('size',0) for p in sh['paragraphs'] for r in p['runs']),default=0))
 titletext=' '.join(pa['text'] for pa in title['paragraphs'])
 for sh in texts:
  sh['role']='title' if sh is title else 'text'
 if i==30:
  # User's explicit final-slide capitalization overrides source casing only.
  title['paragraphs'][0]['runs'][0]['text']='Dinlediğiniz İçin Teşekkürler';title['paragraphs'][0]['text']='Dinlediğiniz İçin Teşekkürler';titletext='Dinlediğiniz İçin Teşekkürler'
  for sh in texts:
   for pa in sh['paragraphs']:
    if pa['text']=='HAZIRLAYANLAR':pa['text']='Hazırlayanlar';pa['runs'][0]['text']='Hazırlayanlar'
 slides.append({'number':i+1,'title':titletext,'background':'#0A0D14','shapes':shapes})
(ROOT/'src/data/slides/slides.json').write_text(json.dumps(slides,ensure_ascii=False,indent=2))
(ROOT/'sources/extraction-manifest.json').write_text(json.dumps({'slides':len(slides),'width':n(p.slide_width),'height':n(p.slide_height),'assets':assets},ensure_ascii=False,indent=2))
print(f'{len(slides)} slides; {len(assets)} unique images; {sum(1 for s in slides for sh in s["shapes"] if sh["type"]=="table")} real tables')
