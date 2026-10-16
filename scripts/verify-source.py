"""Compare all original text and cells against the editable slide data."""
from pptx import Presentation
from pathlib import Path
import json
root=Path(__file__).resolve().parents[1]
p=Presentation(root/'sources/android-ios.pptx');slides=json.loads((root/'src/data/slides/slides.json').read_text(encoding='utf-8'))
count=0
for i,(source,dest) in enumerate(zip(p.slides,slides)):
 for sh in source.shapes:
  target=next((s for s in dest['shapes'] if not s['background'] and s['id']==sh.shape_id),None)
  if sh.has_text_frame and sh.text.strip():
   actual='\n'.join(pa['text'] for pa in target['paragraphs'])
   expected=sh.text
   if i==30:expected=expected.replace('Dinlediğiniz için teşekkürler','Dinlediğiniz İçin Teşekkürler').replace('HAZIRLAYANLAR','Hazırlayanlar')
   assert actual==expected,(i+1,sh.shape_id,expected,actual)
   count+=1
  elif sh.has_table:
   for r,row in enumerate(sh.table.rows):
    for c,cell in enumerate(row.cells):
     actual='\n'.join(pa['text'] for pa in target['rows'][r]['cells'][c]['paragraphs'])
     assert actual==cell.text,(i+1,r,c)
     count+=1
print(f'PASS: {count} text shapes / table cells match PPTX. Only two requested final-slide casing changes.')
